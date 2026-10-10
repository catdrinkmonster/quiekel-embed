"""Every kind of file the app reads, made from scratch (see samples.py)."""

import zipfile
from pathlib import Path

import pytest

import samples
from quiekel_embed import config, documents, mail
from quiekel_embed.extract import Damaged, SkipFile, extract, has_preview, kind_of, picture, problem


def text_of(path: Path) -> str:
    return "\n".join(c.snippet for c in extract(path, path.stat().st_size).chunks)


def test_libreoffice_text_with_its_spaces_tables_and_footnotes(tmp_path):
    odt = samples.opendocument(tmp_path / "Mietvertrag.odt", (
        '<text:h>Mietvertrag</text:h>'
        '<text:p>Die Miete beträgt<text:s text:c="2"/>800 Euro<text:tab/>monatlich.'
        '<text:note><text:note-body><text:p>Ohne Nebenkosten</text:p></text:note-body></text:note></text:p>'
        '<text:tracked-changes><text:changed-region><text:deletion><text:p>alter Betrag</text:p>'
        '</text:deletion></text:changed-region></text:tracked-changes>'
        '<table:table><table:table-row><table:table-cell><text:p>Kaution</text:p></table:table-cell>'
        '<table:table-cell><text:p>2400</text:p></table:table-cell></table:table-row></table:table>'))
    text = text_of(odt)
    assert "Die Miete beträgt 800 Euro monatlich." in text  # (runs of spaces collapse when chunked)
    assert "Ohne Nebenkosten" in text and "Kaution | 2400" in text
    assert "alter Betrag" not in text  # deleted in tracked changes


def test_libreoffice_spreadsheets_and_presentations(tmp_path):
    ods = samples.opendocument(tmp_path / "Haushalt.ods", (
        '<table:table table:name="März"><table:table-row><table:table-cell><text:p>Strom</text:p>'
        '</table:table-cell><table:table-cell><text:p>59,50</text:p></table:table-cell></table:table-row>'
        '<table:table-row table:number-rows-repeated="1048570"><table:table-cell/></table:table-row></table:table>'),
        kind="spreadsheet")
    assert "Sheet März:\nStrom | 59,50" in documents.read(ods)
    odp = samples.opendocument(tmp_path / "Urlaub.odp", (
        '<draw:page draw:name="page1"><draw:frame><draw:text-box><text:p>Italien 2026</text:p>'
        '</draw:text-box></draw:frame></draw:page>'), kind="presentation")
    assert "Slide 1:\nItalien 2026" in documents.read(odp)


def test_libreoffice_documents_with_a_password_are_skipped(tmp_path):
    odt = samples.opendocument(tmp_path / "geheim.odt", "<text:p>x</text:p>", encrypted=True)
    with pytest.raises(SkipFile, match="password"):
        extract(odt, odt.stat().st_size)


def test_flat_opendocument(tmp_path):
    fodt = tmp_path / "notiz.fodt"
    fodt.write_text(f'<?xml version="1.0"?><office:document {samples._ODF_NS}><office:body><office:text>'
                    '<text:p>Flache Datei</text:p></office:text></office:body></office:document>', encoding="utf-8")
    assert text_of(fodt) == "Flache Datei"


def test_libreoffice_documents_show_their_own_preview(tmp_path):
    odt = samples.opendocument(tmp_path / "mit_bild.odt", "<text:p>x</text:p>", thumbnail=samples.png((256, 180)))
    assert has_preview(odt.name, "doc")
    assert picture(odt, 360).size == (256, 180)


def _with_content_type(src: Path, dest: Path, old: bytes, new: bytes) -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item)
            zout.writestr(item, data.replace(old, new) if item.filename == "[Content_Types].xml" else data)
    return dest


def test_word_macro_documents_and_templates(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Angebot Dachsanierung")
    d.save(tmp_path / "a.docx")
    main = b"wordprocessingml.document.main+xml"
    for ext, ctype in ((".docm", b"application/vnd.ms-word.document.macroEnabled.main+xml"),
                       (".dotx", b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml")):
        variant = _with_content_type(tmp_path / "a.docx", tmp_path / f"a{ext}",
                                     b"application/vnd.openxmlformats-officedocument." + main, ctype)
        assert text_of(variant) == "Angebot Dachsanierung"


def test_powerpoint_slide_shows_with_their_notes(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Urlaubsplanung"
    slide.notes_slide.notes_text_frame.text = "Hund nicht vergessen"
    prs.save(tmp_path / "u.pptx")
    show = _with_content_type(tmp_path / "u.pptx", tmp_path / "u.ppsx", b"presentation.main+xml", b"slideshow.main+xml")
    text = text_of(show)
    assert "Urlaubsplanung" in text and "Notes: Hund nicht vergessen" in text


def test_word_97_with_its_fields(tmp_path):
    doc = tmp_path / "Brief.doc"
    doc.write_bytes(samples.word_97(
        'Sehr geehrte Frau Beispiel,\rdie Miete steigt \x13 HYPERLINK "https://example.com" \x14ab Mai\x15.\r'))
    assert text_of(doc) == "Sehr geehrte Frau Beispiel,\ndie Miete steigt ab Mai."


def test_excel_97(tmp_path):
    xls = tmp_path / "Haushalt.xls"
    xls.write_bytes(samples.excel_97("Kosten", [("Miete", 800), ("Strom", 59.5)]))
    assert text_of(xls) == "Sheet Kosten:\nMiete | 800\nStrom | 59.5"


def test_powerpoint_97_without_its_master_slides(tmp_path):
    ppt = tmp_path / "Vortrag.ppt"
    ppt.write_bytes(samples.powerpoint_97(["Urlaubsplanung", "Strand und Berge"]))
    text = text_of(ppt)
    assert "Slide 1:\nUrlaubsplanung" in text and "Slide 2:\nStrand und Berge" in text
    assert "Master" not in text


def test_an_old_office_name_with_something_else_inside(tmp_path):
    rtf = tmp_path / "Brief.doc"  # Word saves RTF like this too
    rtf.write_bytes(rb"{\rtf1\ansi Hallo \'e4\par Welt}")
    assert text_of(rtf) == "Hallo ä\nWelt"
    page = tmp_path / "Tabelle.xls"  # Excel's "web page" files
    page.write_text("<html><body><table><tr><td>Umsatz</td></tr></table></body></html>", encoding="utf-8")
    assert text_of(page) == "Umsatz"


def test_rtf(tmp_path):
    rtf = tmp_path / "Brief.rtf"
    rtf.write_bytes(
        rb"{\rtf1\ansi\ansicpg1252{\fonttbl{\f0 Arial;}}{\colortbl;\red0\green0\blue0;}"
        rb"{\*\generator Riched20;}\f0 Gr\'fc\'dfe aus K\'f6ln\par Preis: 5" + rb"\u" + rb"8364?\tab netto"
        rb"{\field{\*\fldinst HYPERLINK x}{\fldrslt Link}}\line Ende}")
    assert text_of(rtf) == "Grüße aus Köln\nPreis: 5€ nettoLink\nEnde"


def test_emails_saved_from_outlook(tmp_path):
    msg = tmp_path / "Rechnung.msg"
    msg.write_bytes(samples.outlook_msg("Ihre Rechnung für März", "Anbei die Rechnung.",
                                        attachments=[("rechnung.txt", b"Betrag 59,50 Euro " * 300)]))
    text = text_of(msg)
    assert text.startswith("Subject: Ihre Rechnung für März\nFrom: Stadtwerke <info@stadtwerke.example>")
    assert "To: Max Muster" in text and "Date: 2026-03-03" in text
    assert "Attachments: rechnung.txt" in text and text.endswith("Anbei die Rechnung.")


def test_a_msg_file_that_isnt_an_email(tmp_path):
    catalog = tmp_path / "de.msg"  # other programs use .msg for their own text files
    catalog.write_text("::msgcat::mcset de Hello Hallo", encoding="utf-8")
    assert text_of(catalog) == "::msgcat::mcset de Hello Hallo"


def test_emails_and_saved_web_pages_in_the_standard_format(tmp_path):
    eml = tmp_path / "Termin.eml"
    eml.write_bytes(b"From: Praxis <praxis@example.com>\r\nTo: max@example.com\r\n"
                    b"Subject: =?utf-8?q?Terminbest=C3=A4tigung?=\r\nDate: Tue, 03 Mar 2026 10:15:00 +0100\r\n"
                    b"Content-Type: text/html; charset=utf-8\r\n\r\n<p>Ihr Termin am <b>5. M\xc3\xa4rz</b></p>")
    text = text_of(eml)
    assert text.startswith("Subject: Terminbestätigung\nFrom: Praxis <praxis@example.com>")
    assert "Ihr Termin am 5. März" in text
    mht = tmp_path / "Seite.mht"
    mht.write_bytes(b"Subject: Rezept Apfelkuchen\r\nContent-Type: multipart/related; boundary=X\r\n\r\n"
                    b"--X\r\nContent-Type: text/html\r\n\r\n<h1>Apfelkuchen</h1><p>Zimt nicht vergessen</p>\r\n"
                    b"--X\r\nContent-Type: image/png\r\nContent-Location: bild.png\r\n\r\nPNG\r\n--X--\r\n")
    text = text_of(mht)
    assert "Rezept Apfelkuchen" in text and "Zimt nicht vergessen" in text and "Attachments" not in text


def test_compressed_rtf_bodies_of_old_outlook_emails():
    plain = b"{\\rtf1 Hallo}"
    stored = (len(plain) + 12).to_bytes(4, "little") + len(plain).to_bytes(4, "little") + b"MELA" + bytes(4) + plain
    assert mail.decompress_rtf(stored) == plain
    # compressed: a reference to "{\rtf1" in the starting dictionary, then " Hallo}" as is
    body = bytes([0b00000001]) + (0 << 4 | (6 - 2)).to_bytes(2, "big") + b" Hallo}"
    packed = (len(body) + 12).to_bytes(4, "little") + (13).to_bytes(4, "little") + b"LZFu" + bytes(4) + body
    assert mail.decompress_rtf(packed) == b"{\\rtf1 Hallo}"


def test_ebooks(tmp_path):
    book = samples.epub(tmp_path / "Roman.epub", "Der Roman", "Es war einmal ein Schwein.")
    assert "Es war einmal ein Schwein." in text_of(book)
    assert has_preview(book.name, "doc")


def test_camera_raw_files_are_seen_by_their_preview(tmp_path):
    raw = tmp_path / "DSC_0042.NEF"
    raw.write_bytes(samples.camera_raw((600, 400)))
    ex = extract(raw, raw.stat().st_size)
    assert ex.kind == "image" and ex.chunks[0].image.size == (600, 400)


def test_svg_pictures_and_icons(tmp_path):
    svg = tmp_path / "plan.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="300" height="200">'
                   '<rect width="300" height="200" fill="#f54e00"/></svg>', encoding="utf-8")
    width, height = extract(svg, svg.stat().st_size).chunks[0].image.size
    assert width == 1024 and abs(height - 683) <= 1
    icon = tmp_path / "icon.svg"
    icon.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24"/>', encoding="utf-8")
    with pytest.raises(SkipFile, match="tiny_image"):
        extract(icon, icon.stat().st_size)


def test_text_in_utf16_and_files_known_by_name(tmp_path):
    reg = tmp_path / "export.reg"
    reg.write_bytes("Windows Registry Editor Version 5.00\r\n[HKEY_CURRENT_USER\\Test]".encode("utf-16"))
    assert text_of(reg).startswith("Windows Registry Editor")
    assert kind_of(Path("Makefile")) == "code" and kind_of(Path("README")) == "text"
    assert kind_of(Path("notes")) is None


def test_every_type_in_the_lists_has_a_reader():
    for ext in config.DOC_EXTS | config.TEXT_EXTS | config.IMAGE_EXTS:
        assert kind_of(Path(f"file{ext}")) is not None


def test_broken_files_are_named_in_words(tmp_path):
    bad = tmp_path / "kaputt.doc"
    bad.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 not really")
    with pytest.raises(Exception) as e:
        extract(bad, bad.stat().st_size)
    assert problem(e.value) == "!damaged_document"
    assert problem(Damaged("email")) == "!damaged_email"
    locked = PermissionError(13, "The process cannot access the file", None, 32)
    assert problem(locked) == "!in_use"
