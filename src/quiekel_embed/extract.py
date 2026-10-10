"""Turn files into embeddable chunks: text passages or images."""

import html
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageFile, ImageOps, UnidentifiedImageError

from . import config

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except ImportError:  # HEIC photos are skipped without it
    pass

Image.MAX_IMAGE_PIXELS = 250_000_000
# A photo missing its last few bytes (an interrupted copy) still shows nearly everything.
ImageFile.LOAD_TRUNCATED_IMAGES = True
_QUOTED_PATH = re.compile(r"""['"][A-Za-z]:[\\/][^'"]*['"]""")


class SkipFile(Exception):
    """The file is fine but has nothing worth indexing."""


@dataclass
class Chunk:
    snippet: str  # shown in search results
    text: str | None = None  # embedded as a document passage
    image: Image.Image | None = None  # embedded as an image


@dataclass
class Extracted:
    kind: str
    chunks: list[Chunk] = field(default_factory=list)


def kind_of(path: Path) -> str | None:
    ext = path.suffix.lower()
    if ext in config.IMAGE_EXTS:
        return "image"
    if ext in config.DOC_EXTS:
        return "doc"
    if ext in config.CODE_EXTS:
        return "code"
    if ext in config.TEXT_EXTS:
        return "text"
    return None


def problem(e: Exception) -> str:
    """Why a file couldn't be read: "!key" where it's a known case (the app puts it in words),
    else the error itself, without file paths."""
    import zipfile

    name = type(e).__name__
    if isinstance(e, PermissionError):
        return "!no_access"
    if isinstance(e, UnidentifiedImageError) or "image file is truncated" in str(e):
        return "!damaged_image"
    if name in ("FileDataError", "EmptyFileError"):  # PyMuPDF
        return "!damaged_pdf"
    if isinstance(e, (zipfile.BadZipFile, KeyError)):  # Word, PowerPoint and Excel files are ZIPs inside
        return "!damaged_document"
    return f"{name}: {_QUOTED_PATH.sub('', str(e))}".strip()[:300]


def max_bytes(kind: str | None) -> int:
    """The biggest file of this kind that gets indexed."""
    return {"image": config.MAX_IMAGE_BYTES, "doc": config.MAX_DOC_BYTES}.get(kind, config.MAX_TEXT_BYTES)


def extract(path: Path, size: int, title: str | None = None) -> Extracted:
    """title: what the model sees as the document title (defaults to the file name)."""
    kind = kind_of(path)
    ext = path.suffix.lower()
    if size == 0:
        raise SkipFile("empty file")
    if kind == "image":
        if size > config.MAX_IMAGE_BYTES:
            raise SkipFile("image too large")
        try:
            image = load_image(path)
        except UnidentifiedImageError:
            with path.open("rb") as f:
                if f.read(4) == b"PK\x03\x04":  # an archive with a picture's name, e.g. an animated sticker
                    raise SkipFile("an archive, not an image") from None
            raise
        return Extracted(kind, [Chunk(snippet="", image=image)])

    if kind == "doc":
        if size > config.MAX_DOC_BYTES:
            raise SkipFile("document too large")
        if ext == ".pdf":
            text, pages = read_pdf(path)
            if not text.strip() and pages:
                return Extracted(kind, [
                    Chunk(snippet=f"Page {i + 1} (scanned)", image=img)
                    for i, img in enumerate(pages)
                ])
        elif ext == ".docx":
            text = read_docx(path)
        elif ext == ".pptx":
            text = read_pptx(path)
        else:
            text = read_xlsx(path)
    else:
        if size > config.MAX_TEXT_BYTES:
            raise SkipFile("text file too large")
        text = read_text(path)
        if ext == ".ipynb":
            text = notebook_text(text)
        elif ext in {".html", ".htm"}:
            text = strip_html(text)

    title = title or path.name
    chunks = [
        Chunk(snippet=c, text=f"title: {title} | text: {c}") for c in chunk_text(text)
    ]
    if not chunks:
        raise SkipFile("no text")
    return Extracted(kind, chunks)


# ---- readers -------------------------------------------------------------


def read_text(path: Path) -> str:
    data = path.read_bytes()
    if b"\x00" in data[:8192]:
        raise SkipFile("binary file")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def notebook_text(raw: str) -> str:
    try:
        nb = json.loads(raw)
    except json.JSONDecodeError:
        return raw
    parts = []
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        parts.append("".join(src) if isinstance(src, list) else str(src))
    return "\n\n".join(parts)


_SCRIPT_RE = re.compile(r"<(script|style)\b.*?</\1>", re.S | re.I)
_TAG_RE = re.compile(r"<[^>]+>")


def strip_html(raw: str) -> str:
    return html.unescape(_TAG_RE.sub(" ", _SCRIPT_RE.sub(" ", raw)))


def read_pdf(path: Path) -> tuple[str, list[Image.Image]]:
    import pymupdf

    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise SkipFile("password protected")
        parts = []
        for page in doc:
            parts.append(page.get_text())
            if sum(len(p) for p in parts) > config.CHUNK_CHARS * config.MAX_CHUNKS_PER_FILE:
                break
        text = "\n\n".join(parts)
        pages = []
        if not text.strip():
            for page in doc.pages(0, min(config.SCANNED_PDF_PAGES, doc.page_count)):
                pix = page.get_pixmap(dpi=100)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                img.thumbnail((config.MAX_IMAGE_SIDE, config.MAX_IMAGE_SIDE))
                pages.append(img)
    return text, pages


def read_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def read_pptx(path: Path) -> str:
    from pptx import Presentation

    parts = []
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        parts.append(f"\n\nSlide {n}:")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def read_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts, rows = [], 0
    try:
        for ws in wb.worksheets:
            parts.append(f"\n\nSheet {ws.title}:")
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    parts.append(" | ".join(cells))
                    rows += 1
                if rows > 5000:
                    return "\n".join(parts)
    finally:
        wb.close()
    return "\n".join(parts)


def load_image(path: Path, max_side: int = config.MAX_IMAGE_SIDE) -> Image.Image:
    with Image.open(path) as src:
        if min(src.size) < config.MIN_IMAGE_SIDE:
            raise SkipFile("image too small")
        # JPEG can decode at reduced size directly, which is much faster for big photos.
        src.draft("RGB", (max_side * 2, max_side * 2))
        img = ImageOps.exif_transpose(src)
        if img.mode in ("RGBA", "LA", "P", "PA"):
            img = img.convert("RGBA")
            bg = Image.new("RGB", img.size, (255, 255, 255))
            bg.paste(img, mask=img.getchannel("A"))
            img = bg
        else:
            img = img.convert("RGB")
    img.thumbnail((max_side, max_side))
    return img


# ---- chunking ------------------------------------------------------------

_BLANK_LINES_RE = re.compile(r"\n{3,}")
_TRAILING_RE = re.compile(r"[^\S\n]+\n")
# Collapse runs of spaces inside a line, but keep leading indentation (code).
_INNER_SPACES_RE = re.compile(r"(?<=\S)[^\S\n]{2,}")


def chunk_text(text: str) -> list[str]:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ")
    text = _TRAILING_RE.sub("\n", text)
    text = _INNER_SPACES_RE.sub(" ", text)
    text = _BLANK_LINES_RE.sub("\n\n", text).strip()
    if not text:
        return []
    size, overlap = config.CHUNK_CHARS, config.CHUNK_OVERLAP
    chunks, start, n = [], 0, len(text)
    while start < n and len(chunks) < config.MAX_CHUNKS_PER_FILE:
        end = min(start + size, n)
        if end < n:
            # Break at the strongest boundary in the back half of the window.
            window = text[start:end]
            for sep in ("\n\n", "\n", ". ", " "):
                cut = window.rfind(sep)
                if cut > size // 2:
                    end = start + cut + len(sep)
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = end - overlap
        space = text.find(" ", start, end)
        if space != -1:
            start = space + 1
    return chunks
